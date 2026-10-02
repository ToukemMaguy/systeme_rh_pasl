"""Rapport d'import au format Excel (une feuille par rubrique), pour la RH."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def rapport_en_excel(rapport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    for ligne in rapport.resume().split("\n"):
        ws.append([ligne])
    ws["A1"].font = Font(bold=True, size=12, color="06547A")
    ws.column_dimensions["A"].width = 110

    def feuille(titre, entetes, lignes, largeurs):
        if not lignes:
            return
        f = wb.create_sheet(titre[:31])
        f.append(entetes)
        for c in f[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="06547A")
        for l in lignes:
            f.append(list(l) if isinstance(l, (list, tuple)) else [l])
        for i, w in enumerate(largeurs):
            f.column_dimensions[chr(65 + i)].width = w
        f.freeze_panes = "A2"

    # Emplacement exact de chaque point dans le fichier de la RH : dans Excel, Ctrl+G puis taper la cellule
    def _ordre(a):
        chiffres = "".join(ch for ch in a["cellule"].split("/")[0] if ch.isdigit())
        return (a["gravite"] != "À corriger", a["feuille"], int(chiffres or 0))
    anomalies = sorted(rapport.anomalies, key=_ordre)
    feuille("À corriger dans le fichier", ["Gravité", "Feuille", "Cellule (Ctrl+G)", "Employé", "Problème", "Que faire"],
            [(a["gravite"], a["feuille"], a["cellule"], a["nom"], a["probleme"], a["action"]) for a in anomalies],
            [14, 26, 16, 38, 70, 60])
    autres = [t for t in rapport.a_verifier if not t.startswith("[")]
    feuille("Autres points", ["Point à vérifier par la RH"], autres, [150])
    feuille("Écritures inconnues", ["Écriture non reconnue (à ajouter au référentiel)"], rapport.inconnus, [150])
    feuille("Créés", ["Employé", "Poste et lieu"], rapport.crees, [45, 70])
    feuille("Modifiés", ["Employé", "Changements"], [(n, " | ".join(c)) for n, c in rapport.modifies], [45, 120])
    feuille("N+1 pré-remplis", ["Employé", "N+1 (chef d'agence)"], rapport.n_plus_1, [45, 45])
    feuille("Mis à disposition", ["Ajoutés au registre"], rapport.mad_crees, [60])
    feuille("Absents du fichier", ["Employé actif dans l'application mais absent du fichier (aucune action)"], rapport.absents_du_fichier, [110])
    feuille("Pour information", ["Différence sans conséquence"], rapport.infos, [150])
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
