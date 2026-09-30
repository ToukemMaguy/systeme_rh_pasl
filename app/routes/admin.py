"""Administration : comptes utilisateurs et référentiels."""
from datetime import datetime, timedelta
from urllib.parse import quote
from fastapi import APIRouter, Request, Depends, Form, HTTPException
from io import BytesIO
from fastapi.responses import RedirectResponse, Response
from sqlalchemy import func
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from ..database import get_db
from .. import models
from .. import auth
from ..outils import _date_ou_none, _entier_ou_none, _exiger_admin, templates
from ..journal import OBJETS_SUIVIS, LIBELLES_COLONNES
from ..securite import valider_nouveau_mot_de_passe, generer_mot_de_passe_provisoire, LONGUEUR_MIN_MDP

router = APIRouter()


# ============================================================
# GESTION DES UTILISATEURS
# ============================================================

ROLES_VALIDES = ("Administrateur", "Chef RH", "Assistant RH", "Comité de direction", "Employé")


def _page_utilisateurs(request: Request, db: Session, mot_de_passe_provisoire: dict | None = None):
    """Affiche la page des comptes. mot_de_passe_provisoire = {"email": ..., "mdp": ...} : affiché une
    seule fois à l'administrateur après une réinitialisation (jamais dans l'adresse de la page)."""
    utilisateurs = (
        db.query(models.Utilisateur)
        .order_by(models.Utilisateur.nom)
        .all()
    )
    employes_actifs = (
        db.query(models.Employe)
        .filter_by(statut="actif")
        .order_by(models.Employe.nom)
        .all()
    )
    employes_lies_ids = {u.employe_id for u in utilisateurs if u.employe_id}
    employes_disponibles = [e for e in employes_actifs if e.id not in employes_lies_ids]

    return templates.TemplateResponse(request, "admin_utilisateurs.html", {
        "utilisateurs": utilisateurs,
        "employes_disponibles": employes_disponibles,
        "roles": ROLES_VALIDES,
        "maintenant": datetime.now(),
        "longueur_min": LONGUEUR_MIN_MDP,
        "mot_de_passe_provisoire": mot_de_passe_provisoire,
        # [PRODUCTION - POINT 3] employés actifs sans compte (création en une fois)
        "nb_sans_compte": len(employes_disponibles),
    })


# ------------------------------------------------------------------
# [PRODUCTION - POINT 3] Création en une fois des comptes de tous les employés actifs
# ------------------------------------------------------------------
def _choisir_identifiant(employe, deja_pris: set[str]) -> str | None:
    """Identifiant de connexion : email pro, sinon email perso, sinon matricule (décision PASL).
    Un identifiant déjà utilisé par un autre compte est sauté."""
    for candidat in (employe.email_pro, employe.email_perso, employe.matricule):
        candidat = (candidat or "").strip()
        if candidat and candidat.lower() not in deja_pris:
            return candidat
    return None


@router.post("/admin/utilisateurs/creer-comptes-manquants")
def creer_comptes_manquants(request: Request, db: Session = Depends(get_db)):
    """Crée un compte « Employé » pour chaque employé actif qui n'en a pas, puis renvoie un fichier Excel
    avec les identifiants et mots de passe provisoires à distribuer. Les mots de passe ne sont enregistrés
    nulle part en clair : ce fichier est le SEUL endroit où ils figurent."""
    _exiger_admin(request)
    deja_pris = {e.lower() for (e,) in db.query(models.Utilisateur.email).all()}
    avec_compte = {i for (i,) in db.query(models.Utilisateur.employe_id).filter(models.Utilisateur.employe_id.isnot(None)).all()}
    employes = (
        db.query(models.Employe)
        .filter(models.Employe.statut == "actif")
        .order_by(models.Employe.nom, models.Employe.prenom)
        .all()
    )
    crees, ignores = [], []
    for e in employes:
        if e.id in avec_compte:
            continue
        identifiant = _choisir_identifiant(e, deja_pris)
        if not identifiant:
            ignores.append((e, "Ni email ni matricule disponible (ou déjà utilisés par un autre compte)"))
            continue
        mdp = generer_mot_de_passe_provisoire()
        db.add(models.Utilisateur(
            email=identifiant, nom=f"{e.prenom} {e.nom}", role="Employé", employe_id=e.id,
            mot_de_passe_hash=auth.hash_password(mdp), doit_changer_mdp=True,
        ))
        deja_pris.add(identifiant.lower())
        crees.append((e, identifiant, mdp))
    db.commit()

    if not crees and not ignores:
        return RedirectResponse(url=f"/admin/utilisateurs?ok={quote('Tous les employés actifs ont déjà un compte.')}", status_code=303)
    if not crees:
        noms = ", ".join(f"{e.nom} {e.prenom}" for e, _ in ignores[:10])
        message = (f"Aucun compte créé : {len(ignores)} employé(s) n'ont ni email ni matricule utilisable ({noms}). "
                   "Complétez leur fiche puis recommencez.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote(message)}", status_code=303)
    return _excel_comptes(crees, ignores)


def _excel_comptes(crees, ignores) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Comptes créés"
    ws.append([f"Comptes créés le {datetime.now().strftime('%d/%m/%Y à %H:%M')} — DOCUMENT CONFIDENTIEL : "
               "à imprimer et distribuer, puis à SUPPRIMER (il contient des mots de passe)."])
    ws["A1"].font = Font(bold=True, color="9B2C2C")
    ws.append(["Adresse de l'application : voir l'intranet / le service RH. Connexion avec l'identifiant ci-dessous "
               "(ou le matricule). Le mot de passe pourra être changé depuis l'icône 🔑 du menu."])
    ws.append([])
    entetes = ["Matricule", "Nom", "Prénom", "Agence / Département", "Identifiant de connexion", "Mot de passe provisoire"]
    ws.append(entetes)
    for cellule in ws[4]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="06547A")
    for e, identifiant, mdp in crees:
        unite = e.agence.nom if e.agence else (e.departement.nom if e.departement else "")
        ws.append([e.matricule or "", e.nom, e.prenom, unite, identifiant, mdp])
    for colonne, largeur in zip("ABCDEF", (14, 22, 22, 28, 34, 24)):
        ws.column_dimensions[colonne].width = largeur
    for ligne in ws.iter_rows(min_row=5, min_col=6, max_col=6):
        for cellule in ligne:
            cellule.font = Font(name="Consolas", size=12, bold=True)
    ws.freeze_panes = "A5"
    if ignores:
        ws2 = wb.create_sheet("Non créés")
        ws2.append(["Matricule", "Nom", "Prénom", "Raison"])
        for e, raison in ignores:
            ws2.append([e.matricule or "", e.nom, e.prenom, raison])
        ws2.column_dimensions["D"].width = 70
    tampon = BytesIO()
    wb.save(tampon)
    nom = f"comptes_crees_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(content=tampon.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nom}"', "Cache-Control": "no-store"})


@router.get("/admin/utilisateurs")
def page_utilisateurs(request: Request, db: Session = Depends(get_db)):
    _exiger_admin(request)
    return _page_utilisateurs(request, db)


@router.post("/admin/utilisateurs/nouveau")
def creer_utilisateur(
    request: Request,
    db: Session = Depends(get_db),
    email: str = Form(...),
    nom: str = Form(...),
    mot_de_passe: str = Form(...),
    role: str = Form(...),
    employe_id: str = Form(""),
):
    _exiger_admin(request)
    if role not in ROLES_VALIDES:
        message = quote("Rôle invalide.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    # [LOT 1 - POINT 3] Même règle de mot de passe partout (8 caractères, lettres et chiffres)
    erreur_mdp = valider_nouveau_mot_de_passe(mot_de_passe)
    if erreur_mdp:
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote(erreur_mdp)}", status_code=303)

    existant = db.query(models.Utilisateur).filter_by(email=email).first()
    if existant:
        message = quote(f"Un compte existe déjà avec l'email {email}.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    if role == "Employé" and not employe_id:
        message = quote("Un compte Employé doit être rattaché à une fiche employé.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    utilisateur = models.Utilisateur(
        email=email,
        nom=nom,
        mot_de_passe_hash=auth.hash_password(mot_de_passe),
        role=role,
        employe_id=_entier_ou_none(employe_id),  # [CORRECTIF B11]
        # [LOT 1 - POINT 3] mot de passe attribué par l'administrateur : signalé comme « provisoire »
        # dans la liste des comptes tant que l'utilisateur ne l'a pas changé (changement NON imposé)
        doit_changer_mdp=True,
    )
    db.add(utilisateur)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        message = quote("Erreur : l'email ou l'employé est déjà utilisé.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    message = quote(f"Compte « {email} » créé.")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


@router.post("/admin/utilisateurs/{user_id}/supprimer")
def supprimer_utilisateur(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    _exiger_admin(request)

    if user_id == request.session.get("user_id"):
        message = quote("Tu ne peux pas supprimer ton propre compte.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        message = quote("Compte introuvable.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    email_supprime = utilisateur.email
    db.delete(utilisateur)
    db.commit()
    message = quote(f"Compte « {email_supprime} » supprimé.")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


@router.post("/admin/utilisateurs/{user_id}/reinitialiser")
def reinitialiser_mot_de_passe(user_id: int, request: Request, db: Session = Depends(get_db)):
    """[LOT 1 - POINT 3] Mot de passe oublié : l'administrateur génère un mot de passe provisoire,
    à transmettre à l'utilisateur, qui pourra le changer quand il le souhaite. Débloque aussi le compte."""
    _exiger_admin(request)
    if user_id == request.session.get("user_id"):
        message = quote("Pour votre propre compte, utilisez « Changer mon mot de passe ».")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)
    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote('Compte introuvable.')}", status_code=303)

    provisoire = generer_mot_de_passe_provisoire()
    utilisateur.mot_de_passe_hash = auth.hash_password(provisoire)
    utilisateur.doit_changer_mdp = True
    utilisateur.tentatives_echouees = 0
    utilisateur.bloque_jusqua = None
    db.commit()
    return _page_utilisateurs(request, db, {"email": utilisateur.email, "mdp": provisoire})


@router.post("/admin/utilisateurs/{user_id}/debloquer")
def debloquer_utilisateur(user_id: int, request: Request, db: Session = Depends(get_db)):
    """[LOT 1 - POINT 3] Lève le blocage temporaire après trop de tentatives de connexion."""
    _exiger_admin(request)
    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote('Compte introuvable.')}", status_code=303)
    utilisateur.tentatives_echouees = 0
    utilisateur.bloque_jusqua = None
    db.commit()
    message = quote(f"Compte « {utilisateur.email} » débloqué.")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


@router.post("/admin/utilisateurs/{user_id}/desactiver")
def desactiver_utilisateur(user_id: int, request: Request, db: Session = Depends(get_db)):
    """[LOT 1 - POINT 4] Coupe l'accès d'un compte sans le supprimer (suspension, erreur, départ anticipé)."""
    _exiger_admin(request)
    if user_id == request.session.get("user_id"):
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote('Vous ne pouvez pas désactiver votre propre compte.')}", status_code=303)
    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote('Compte introuvable.')}", status_code=303)
    utilisateur.compte_actif = False
    utilisateur.desactive_le = datetime.now()
    utilisateur.motif_desactivation = f"Désactivé par l'administrateur ({request.session.get('nom')})"
    db.commit()
    message = quote(f"Compte « {utilisateur.email} » désactivé : l'utilisateur est déconnecté immédiatement.")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


@router.post("/admin/utilisateurs/{user_id}/reactiver")
def reactiver_utilisateur(user_id: int, request: Request, db: Session = Depends(get_db)):
    """[LOT 1 - POINT 4] Rétablit l'accès, sauf si l'employé rattaché est sorti des effectifs."""
    _exiger_admin(request)
    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={quote('Compte introuvable.')}", status_code=303)
    if utilisateur.employe and utilisateur.employe.statut != "actif":
        message = quote("Impossible : l'employé rattaché à ce compte ne fait plus partie des effectifs.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)
    utilisateur.compte_actif = True
    utilisateur.desactive_le = None
    utilisateur.motif_desactivation = None
    utilisateur.tentatives_echouees = 0
    utilisateur.bloque_jusqua = None
    db.commit()
    message = quote(f"Compte « {utilisateur.email} » réactivé.")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


@router.post("/admin/utilisateurs/{user_id}/role")
def changer_role_utilisateur(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
    nouveau_role: str = Form(...),
):
    _exiger_admin(request)

    if nouveau_role not in ROLES_VALIDES:
        message = quote("Rôle invalide.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    if user_id == request.session.get("user_id"):
        message = quote("Tu ne peux pas modifier ton propre rôle.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    utilisateur = db.query(models.Utilisateur).get(user_id)
    if not utilisateur:
        message = quote("Compte introuvable.")
        return RedirectResponse(url=f"/admin/utilisateurs?erreur={message}", status_code=303)

    utilisateur.role = nouveau_role
    db.commit()
    message = quote(f"Rôle mis à jour pour « {utilisateur.email} ».")
    return RedirectResponse(url=f"/admin/utilisateurs?ok={message}", status_code=303)


# ============================================================
# ADMINISTRATION DES RÉFÉRENTIELS
# ============================================================

REFERENTIELS = {
    "departements": {"modele": models.Departement, "champ": "nom", "label": "Départements", "icone": "fa-building"},
    "agences": {"modele": models.Agence, "champ": "nom", "label": "Agences", "icone": "fa-map-marker-alt"},
    "postes": {"modele": models.Poste, "champ": "intitule", "label": "Postes", "icone": "fa-briefcase"},
    "grades": {"modele": models.Grade, "champ": "libelle", "label": "Grades", "icone": "fa-medal"},
}


@router.get("/admin")
def page_admin(request: Request, section: str = "departements", db: Session = Depends(get_db)):
    _exiger_admin(request)
    if section not in REFERENTIELS:
        section = "departements"

    modele = REFERENTIELS[section]["modele"]
    champ = REFERENTIELS[section]["champ"]
    elements = db.query(modele).order_by(getattr(modele, champ)).all()

    compteurs = {}
    champ_fk = {"departements": "departement_id", "agences": "agence_id", "postes": "poste_id", "grades": "grade_id"}[section]
    for e in elements:
        compteurs[e.id] = db.query(models.Employe).filter_by(**{champ_fk: e.id}).count()

    return templates.TemplateResponse(request, "admin.html", {
        "section": section,
        "referentiels": REFERENTIELS,
        "elements": elements,
        "compteurs": compteurs,
    })


@router.post("/admin/{section}/nouveau")
def ajouter_referentiel(
    section: str,
    request: Request,
    db: Session = Depends(get_db),
    valeur: str = Form(...),
    ville: str = Form(""),
    region: str = Form(""),
    niveau_hierarchique: str = Form(""),
):
    _exiger_admin(request)
    if section not in REFERENTIELS:
        raise HTTPException(status_code=404, detail="Référentiel inconnu")

    modele = REFERENTIELS[section]["modele"]
    champ = REFERENTIELS[section]["champ"]

    if section == "agences":
        obj = modele(**{champ: valeur, "ville": ville or None, "region": region or None})
    elif section == "grades":
        # [CORRECTIF B11] niveau non numérique → message au lieu d'une erreur 500
        if niveau_hierarchique and _entier_ou_none(niveau_hierarchique) is None:
            message = quote("Le niveau hiérarchique doit être un nombre entier.")
            return RedirectResponse(url=f"/admin?section={section}&erreur={message}", status_code=303)
        obj = modele(**{champ: valeur, "niveau_hierarchique": _entier_ou_none(niveau_hierarchique)})
    else:
        obj = modele(**{champ: valeur})

    db.add(obj)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        message = quote(f"« {valeur} » existe déjà.")
        return RedirectResponse(url=f"/admin?section={section}&erreur={message}", status_code=303)

    return RedirectResponse(url=f"/admin?section={section}", status_code=303)


@router.post("/admin/{section}/{element_id}/supprimer")
def supprimer_referentiel(
    section: str,
    element_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    _exiger_admin(request)
    if section not in REFERENTIELS:
        raise HTTPException(status_code=404, detail="Référentiel inconnu")

    modele = REFERENTIELS[section]["modele"]
    obj = db.query(modele).get(element_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Élément introuvable")

    db.delete(obj)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        message = quote("Impossible de supprimer : des employés y sont encore rattachés.")
        return RedirectResponse(url=f"/admin?section={section}&erreur={message}", status_code=303)

    return RedirectResponse(url=f"/admin?section={section}", status_code=303)


# ============================================================
# JOURNAL D'AUDIT (lot 1, point 1) — consultation seule
# ============================================================

ACTIONS_JOURNAL = {
    "creation": "Création", "modification": "Modification", "suppression": "Suppression",
    "connexion": "Connexion", "echec_connexion": "Échec de connexion", "deconnexion": "Déconnexion",
}
LIGNES_PAR_PAGE = 100


@router.get("/admin/journal")
def page_journal(
    request: Request,
    db: Session = Depends(get_db),
    date_debut: str = "",
    date_fin: str = "",
    utilisateur: str = "",
    objet_type: str = "",
    action: str = "",
    employe: str = "",
    page: int = 1,
):
    """Réservé à l'Administrateur et au Chef RH."""
    if request.session.get("role") not in ("Administrateur", "Chef RH"):
        _exiger_admin(request)

    q = db.query(models.JournalAudit)
    d1, d2 = _date_ou_none(date_debut), _date_ou_none(date_fin)
    if d1:
        q = q.filter(models.JournalAudit.date_heure >= datetime(d1.year, d1.month, d1.day))
    if d2:
        q = q.filter(models.JournalAudit.date_heure < datetime(d2.year, d2.month, d2.day) + timedelta(days=1))
    if utilisateur.strip():
        q = q.filter(models.JournalAudit.utilisateur_nom.ilike(f"%{utilisateur.strip()}%"))
    if objet_type:
        q = q.filter(models.JournalAudit.objet_type == objet_type)
    if action:
        q = q.filter(models.JournalAudit.action == action)
    if employe.strip():
        motif = f"%{employe.strip()}%"
        ids = [e.id for e in db.query(models.Employe.id).filter(
            models.Employe.nom.ilike(motif) | models.Employe.prenom.ilike(motif) | models.Employe.matricule.ilike(motif))]
        q = q.filter(models.JournalAudit.employe_id.in_(ids or [-1]))

    total = q.count()
    page = max(page, 1)
    lignes = (q.order_by(models.JournalAudit.date_heure.desc(), models.JournalAudit.id.desc())
               .offset((page - 1) * LIGNES_PAR_PAGE).limit(LIGNES_PAR_PAGE).all())

    ids_employes = {l.employe_id for l in lignes if l.employe_id}
    noms_employes = {e.id: f"{e.nom} {e.prenom}" for e in
                     db.query(models.Employe).filter(models.Employe.id.in_(ids_employes or [-1]))}

    filtres = {"date_debut": date_debut, "date_fin": date_fin, "utilisateur": utilisateur,
               "objet_type": objet_type, "action": action, "employe": employe}
    return templates.TemplateResponse(request, "journal.html", {
        "lignes": lignes,
        "total": total,
        "page": page,
        "nb_pages": max((total + LIGNES_PAR_PAGE - 1) // LIGNES_PAR_PAGE, 1),
        "filtres": filtres,
        "requete_filtres": "&".join(f"{k}={quote(v)}" for k, v in filtres.items() if v),
        "types_objets": sorted(OBJETS_SUIVIS.values()) + ["Session"],
        "actions": ACTIONS_JOURNAL,
        "libelles": LIBELLES_COLONNES,
        "noms_employes": noms_employes,
    })
