"""Guides utilisateurs à imprimer (PDF).

1) Une page par employé, avec SES accès déjà imprimés (à partir du fichier Excel donné par
   Administration > Utilisateurs > « Créer les comptes manquants ») :

    python -m app.scripts.imprimer_identifiants data\\comptes.xlsx --adresse http://NOM-DU-SERVEUR:8000

   -> data\\acces_personnels_AAAAMMJJ_HHMM.pdf : pages classées par agence / direction puis par nom.
      Un N+1 (qui a des collaborateurs rattachés) reçoit aussi la page « Guide N+1 ».
   Sans --adresse, l'adresse APP_URL du fichier .env est utilisée.

2) Les 4 guides avec l'encadré « Vos accès » à remplir à la main (Employé, N+1, RH, Comité) :

    python -m app.scripts.imprimer_identifiants --vierges

CONFIDENTIEL : le PDF personnalisé contient des mots de passe. Il est écrit dans data\\ (jamais publié sur
GitHub). Imprimez-le, puis supprimez-le ainsi que le fichier Excel des comptes.
"""
import argparse
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

from app.guides_utilisateurs import guides_personnalises, guides_vierges

COLONNES = {"nom": "Nom", "prenom": "Prénom", "unite": "Agence / Département",
            "identifiant": "Identifiant de connexion", "mot_de_passe": "Mot de passe provisoire"}


def lire_comptes(chemin) -> list[dict]:
    """Lit la feuille « Comptes créés » : la ligne d'en-tête est repérée par son contenu, pas par sa position."""
    ws = load_workbook(chemin, read_only=True, data_only=True).worksheets[0]
    lignes = list(ws.iter_rows(values_only=True))
    for n, ligne in enumerate(lignes):
        entetes = [str(v or "").strip() for v in ligne]
        if COLONNES["identifiant"] in entetes and COLONNES["mot_de_passe"] in entetes:
            break
    else:
        raise ValueError("colonnes « Identifiant de connexion » et « Mot de passe provisoire » introuvables : "
                         "utilisez le fichier Excel donné par « Créer les comptes manquants ».")
    position = {cle: entetes.index(titre) for cle, titre in COLONNES.items() if titre in entetes}
    comptes = []
    for ligne in lignes[n + 1:]:
        valeurs = {cle: str(ligne[i] or "").strip() for cle, i in position.items() if i < len(ligne)}
        if valeurs.get("identifiant") and valeurs.get("mot_de_passe"):
            comptes.append(valeurs)
    return comptes


def identifiants_des_n_plus_1() -> set[str] | None:
    """Identifiants (en minuscules) des comptes dont l'employé a au moins un collaborateur actif rattaché."""
    try:
        from app.database import SessionLocal
        from app import models
        db = SessionLocal()
        try:
            chefs = {i for (i,) in db.query(models.Employe.n_plus_1_id)
                     .filter(models.Employe.n_plus_1_id.isnot(None), models.Employe.statut == "actif").all()}
            return {email.lower() for (email, employe_id) in
                    db.query(models.Utilisateur.email, models.Utilisateur.employe_id).all() if employe_id in chefs}
        finally:
            db.close()
    except Exception as e:
        print(f"ATTENTION : base inaccessible ({str(e).splitlines()[0][:150]}) : les pages « Guide N+1 » ne seront pas ajoutées.")
        return None


def main():
    load_dotenv()
    parser = argparse.ArgumentParser(description="Guides utilisateurs PASL-RH à imprimer")
    parser.add_argument("fichier_comptes", nargs="?", help="Excel donné par « Créer les comptes manquants »")
    parser.add_argument("--adresse", default=os.getenv("APP_URL", ""), help="adresse de l'application (ex. http://SERVEUR:8000)")
    parser.add_argument("--vierges", action="store_true", help="les 4 guides avec des accès à remplir à la main")
    args = parser.parse_args()
    Path("data").mkdir(exist_ok=True)

    if args.vierges:
        chemin = Path("data") / "guides_utilisateurs.pdf"
        guides_vierges(chemin)
        print(f"Guides vierges (Employé, N+1, RH, Comité) : {chemin}")
        return
    if not args.fichier_comptes:
        parser.error("indiquez le fichier Excel des comptes, ou --vierges")

    adresse = args.adresse.strip().rstrip("/")
    if not adresse or "127.0.0.1" in adresse or "localhost" in adresse:
        print(f"Adresse de l'application invalide (« {adresse} ») : ajoutez --adresse http://NOM-DU-SERVEUR:8000")
        sys.exit(1)
    try:
        comptes = lire_comptes(args.fichier_comptes)
    except (OSError, ValueError) as e:
        print(f"FICHIER REFUSÉ : {e}")
        sys.exit(1)
    if not comptes:
        print("Aucun compte trouvé dans le fichier.")
        sys.exit(1)

    chefs = identifiants_des_n_plus_1()
    comptes.sort(key=lambda c: (c.get("unite", "").lower(), c.get("nom", "").lower(), c.get("prenom", "").lower()))
    pages, nb_n_plus_1 = [], 0
    for c in comptes:
        acces = {"nom": f"{c.get('nom', '')} {c.get('prenom', '')}".strip(), "unite": c.get("unite", ""),
                 "adresse": adresse, "identifiant": c["identifiant"], "mot_de_passe": c["mot_de_passe"]}
        pages.append(("employe", acces))
        if chefs and c["identifiant"].lower() in chefs:
            pages.append(("n_plus_1", acces))
            nb_n_plus_1 += 1

    chemin = Path("data") / f"acces_personnels_{datetime.now():%Y%m%d_%H%M}.pdf"
    guides_personnalises(chemin, pages)
    print(f"{len(comptes)} employé(s), {len(pages)} page(s) dont {nb_n_plus_1} « Guide N+1 » : {chemin}")
    for unite, nombre in sorted(Counter(c.get("unite") or "(non renseigné)" for c in comptes).items()):
        print(f"   {unite} : {nombre}")
    print("\nCONFIDENTIEL : imprimez ce PDF, puis supprimez-le ainsi que le fichier Excel des comptes.")


if __name__ == "__main__":
    main()
