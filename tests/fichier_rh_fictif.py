"""Fabrique un classeur FICTIF ayant exactement la présentation du fichier RH (fixhiier_r.xlsx),
pour tester l'import sans données personnelles réelles."""
from datetime import datetime
from openpyxl import Workbook


def fabriquer(chemin, variante=None):
    wb = Workbook()
    wb.active.title = "SUIVI Du Personnel 2026"
    ws = wb.create_sheet("Effectif_mois août 2026 ")
    for col, titre in zip("GHIJK", ["Fonction", "Date d'embauche", "Num. tél Prof", "Num. tél personnel", "Adresse Personnelle électronique"]):
        ws[f"{col}2"] = titre
    if variante == "colonne_inseree":
        ws["G2"] = "Matricule"
    lignes = [
        ("DIRECTION GENERALE/ HEAD OFFICE", None),
        (("F", None), ("MBALLA Rose", "Chef Service Comptable", datetime(2015, 1, 5), "699000001", None, "rose@exemple.cm")),
        ("MIS A DISPOSITION", None),
        ((None, "H"), ("OWONA Paul", "Driver", datetime(2022, 3, 1), "677000009", None, None)),
        ("AGENCE DE MVOMGBI", None),
        ((None, "H"), ("ATEBA Jean", "Branch manager", datetime(2012, 6, 1), "699000002", "677000002", None)),
        (("F", None), ("NGONO Alice", "LOI", datetime(2020, 2, 3), "699000003", "677000003", "alice@exemple.cm")),
        (("F", None), ("ESSOMBA Marie", "Cashier Agent P-S", datetime(2021, 9, 1), None, "677000004", None)),
        ("YAOUNDE", None),
        ((None, "H"), ("FOUDA Luc", "Analyste crédt", datetime(2019, 4, 1), None, None, None)),
    ]
    r = 4
    for a, b in lignes:
        if b is None:
            ws.cell(row=r, column=6, value=a)
        else:
            (f, h), (nom, fonction, emb, tp, tperso, mail) = a, b
            ws.cell(row=r, column=2, value=f); ws.cell(row=r, column=4, value=h); ws.cell(row=r, column=5, value=1)
            ws.cell(row=r, column=6, value=nom); ws.cell(row=r, column=7, value=fonction); ws.cell(row=r, column=8, value=emb)
            ws.cell(row=r, column=9, value=tp); ws.cell(row=r, column=10, value=tperso); ws.cell(row=r, column=11, value=mail)
            if variante == "telephone_modifie" and nom == "NGONO Alice":
                ws.cell(row=r, column=9, value="655555555")
            if variante == "telephone_vide" and nom == "NGONO Alice":
                ws.cell(row=r, column=9, value=None)
        r += 1

    ws = wb.create_sheet("Entrées")
    ws["B4"], ws["C4"], ws["D4"] = "MAT", "Noms et Prénoms", "Date d'embauche"
    ws["B5"], ws["C5"], ws["D5"] = 900, "FOUDA LUC", datetime(2019, 4, 1)
    wb.create_sheet("Départs")
    d = wb["Départs"]
    d["B3"] = "Noms et prénoms"
    d.append([1, "ANCIEN Paul", "BSSADI", "LO", None, "DEX", datetime(2010, 1, 1), datetime(2018, 5, 31), "démission"])
    d.append([1, "NGONO ALICE", "CONGO", "CSA", None, "OPERATION", datetime(2012, 1, 1), datetime(2013, 5, 31), "fin de contrat"])
    m = wb.create_sheet("Mouvements internes")
    m["B3"] = "NOMS ET PRENOMS"
    m.append([1, "NGONO Alice", "CSA", "CONGO", "LO", "YAOUNDE / MVOG-MBI", "PROMOTION", datetime(2023, 1, 3)])
    wb.create_sheet("Suivi des absences")
    p = wb.create_sheet("Provision_congés_24j_")
    p["A5"], p["C5"], p["E5"], p["BI5"], p["BP5"] = "MAT", "Noms et Prénoms des employés", "Embauche", "Fin Période", "Total"
    for i, (mat, nom, emb, total, conges) in enumerate([
        (101, "MBALLA ROSE", datetime(2015, 1, 5), 30, [(datetime(2025, 12, 1), datetime(2025, 12, 20))]),
        (102, "ATEBA JEAN", datetime(2012, 6, 1), 12, []),
        (103, "NGONO ALICE", datetime(2020, 2, 3), 18, [(datetime(2026, 7, 6), datetime(2026, 7, 17))]),
        (104, "ESSOMBA MARIE", datetime(2021, 9, 1), 24, []),
    ]):
        l = 6 + i
        p.cell(row=l, column=1, value=mat); p.cell(row=l, column=3, value=nom); p.cell(row=l, column=5, value=emb)
        for j, (d1, d2) in enumerate(conges):
            p.cell(row=l, column=7 + j * 6, value=d1); p.cell(row=l, column=8 + j * 6, value=d2)
        p.cell(row=l, column=61, value=datetime(2026, 8, 31)); p.cell(row=l, column=68, value=total)
        p.cell(row=l, column=69, value=999999)          # salaire : ne doit JAMAIS être lu
    wb.create_sheet("Employés_mis_en_congé_en_2026")
    wb.save(chemin)
